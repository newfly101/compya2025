#!/usr/bin/env python3
"""players.json → 포지션 입력용 엑셀 시트 생성.

인게임 화면과 눈으로 대조하며 포지션을 채워 넣기 위한 작업 시트다.
포지션을 다 채운 뒤에는 이 엑셀을 다시 읽어 players.json 에 합친다.

시트 구성
  선수    카드 1장 = 1행. 포지션을 여기에 채운다
  코치    포지션 개념이 없어 분리. 참고용
  입력값  드롭다운 목록 (포지션 표기 오타 방지)

사용법
  python scripts/gen_player_position_sheet.py [-o 출력경로]
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SOURCE = Path("web/src/data/players/players.json")

# 사용자 지정 12종. 원문의 "CR" 은 중견수 CF 로 본다.
POSITIONS = ["C", "1B", "2B", "3B", "SS", "LF", "CF", "RF", "DH", "SP", "RP", "CP"]
PITCHER_POSITIONS = {"SP", "RP", "CP"}

# 5단 풀세트는 매 행에 늘어놓으면 읽기 어려워 범위로 줄인다.
GRADE_FULL = ["노멀", "레어", "스페셜", "히어로", "플래티넘"]

HEADERS = [
    ("id", 10), ("이름", 12), ("구단", 8), ("연도", 7),
    ("등급", 20), ("카드 종류", 30),
    ("타자/투수", 10), ("포지션", 10), ("비고", 16),
]

HEAD_FILL = PatternFill("solid", fgColor="2A2A3A")
HEAD_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF6D5")  # 사람이 채울 칸


def is_coach(p: dict) -> bool:
    """index.js 의 isCoach 와 같은 규칙 — grades 3종(스페셜·히어로·플래티넘)."""
    g = p.get("grades") or []
    return len(g) == 3 and {"스페셜", "히어로", "플래티넘"} <= set(g)


def grade_text(grades: list[str]) -> str:
    g = list(grades or [])
    if g[: len(GRADE_FULL)] == GRADE_FULL:
        rest = g[len(GRADE_FULL):]
        return "노멀~플래티넘" + (", " + ", ".join(rest) if rest else "")
    return ", ".join(g)


def write_sheet(ws, rows: list[dict], *, with_position: bool) -> None:
    for col, (name, width) in enumerate(HEADERS, start=1):
        c = ws.cell(1, col, name)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    for i, p in enumerate(rows, start=2):
        ws.cell(i, 1, p["id"])
        ws.cell(i, 2, p.get("name", ""))
        ws.cell(i, 3, p.get("team", ""))
        ws.cell(i, 4, p.get("year") if p.get("year") is not None else "레전드")
        ws.cell(i, 5, grade_text(p.get("grades")))
        ws.cell(i, 6, ", ".join(p.get("cardTypes") or []))
        if with_position:
            # 포지션만 채우면 타자/투수는 따라온다 — 손으로 두 번 적지 않게.
            ws.cell(i, 7, f'=IF($H{i}="","",IF(OR($H{i}="SP",$H{i}="RP",$H{i}="CP"),"투수","타자"))')
            ws.cell(i, 8).fill = INPUT_FILL
        for col in (1, 3, 4, 7, 8):
            ws.cell(i, col).alignment = Alignment(horizontal="center")

    last = len(rows) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"
    ws.freeze_panes = "A2"

    if with_position and last > 1:
        dv = DataValidation(
            type="list", formula1='"' + ",".join(POSITIONS) + '"',
            allow_blank=True, showDropDown=False,
            errorTitle="포지션 표기", error="목록에 있는 값만 넣을 수 있다.",
        )
        ws.add_data_validation(dv)
        dv.add(f"H2:H{last}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "-o", "--out", type=Path,
        default=Path("test-docs/레전드 재료 앱 디자인/data_handoff_player_position/선수_포지션_입력.xlsx"),
    )
    args = ap.parse_args()

    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    players = data["players"]

    # 구단 → 연도 → 이름 순. 인게임에서 구단별로 훑어보며 대조하기 좋은 순서다.
    def key(p):
        return (p.get("team") or "", p.get("year") or 0, p.get("name") or "")

    coaches = sorted((p for p in players if is_coach(p)), key=key)
    actives = sorted((p for p in players if not is_coach(p)), key=key)

    wb = Workbook()
    write_sheet(wb.active, actives, with_position=True)
    wb.active.title = "선수"
    write_sheet(wb.create_sheet("코치"), coaches, with_position=False)

    ref = wb.create_sheet("입력값")
    ref["A1"], ref["A1"].font = "포지션", Font(bold=True)
    for i, pos in enumerate(POSITIONS, start=2):
        ref.cell(i, 1, pos)
        ref.cell(i, 2, "투수" if pos in PITCHER_POSITIONS else "타자")
    ref["D1"], ref["D1"].font = "메모", Font(bold=True)
    for i, line in enumerate([
        f"원본: {SOURCE}  ({data.get('source', '')})",
        f"전체 {len(players):,}건 = 선수 {len(actives):,} + 코치 {len(coaches):,}",
        "노란 칸(포지션)만 채우면 타자/투수는 수식으로 따라온다.",
        "id 는 원본과 합칠 때 쓰는 열쇠다 — 지우거나 바꾸지 말 것.",
        "연도 '레전드' 는 원본에 연도가 없는 카드다.",
    ], start=2):
        ref.cell(i, 4, line)
    ref.column_dimensions["D"].width = 70

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    print(f"선수 {len(actives):,}행 / 코치 {len(coaches):,}행")
    print(f"→ {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
