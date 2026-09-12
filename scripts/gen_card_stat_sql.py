# -*- coding: utf-8 -*-
"""노말 카드 스탯 + 구종등급 적재 SQL 생성 — 정본은 게임 화면 실측 엑셀.

card_id 는 엑셀에 없다. 운영 DB 덤프(prod2.tsv, 아래 재현 방법 참고)에서
(구단, 연도, 이름, 역할)로 조회한다. 이름만 쓰면 같은 해에 타자·투수 카드를
둘 다 가진 선수(김성한B, 김정수C 등)가 섞이므로 반드시 역할까지 넣는다.

짝을 못 지은 행이 하나라도 있으면 파일을 만들지 않고 예외로 멈춘다 — 조용히
빼면 나중에 "왜 이 카드만 스탯이 없지"를 다시 조사해야 한다.

prod2.tsv 재현 방법 (운영 DB 에서 1회성으로 뽑은 스냅샷 — SSH/실행 금지 규정상
이 스크립트가 직접 조회하지 않는다):
    SELECT id, player_name, team_code, season_year, position_code,
           sub_position_code, player_role
    FROM data_player_card
    WHERE card_type = 'NORMAL';
    (탭 구분, 헤더 없이 TSV 로 export)
"""
from __future__ import annotations

import io
import uuid

import openpyxl

XLSX = "test-docs/노말선수_스탯,구종등급_정리_최종본.xlsx"
CARD_DUMP = (
    "C:/Users/hibee/AppData/Local/Temp/claude/D--NewProjects-com2usbaseball/"
    "f0df1c25-df82-4082-875e-993491878b67/scratchpad/prod2.tsv"
)
OUT = "sql/V3/data/data_player_card_stat_INSERT.sql"

# 엑셀 한글 구단 표기 → DB team_code (영문 표기는 이미 DB 코드와 같아 그대로 둔다)
TEAM = {
    "두산": "DOO", "롯데": "LOT", "빙그레": "BIN", "삼미": "SUP", "삼성": "SAM",
    "쌍방울": "SSA", "청보": "CHU", "키움": "KIW", "태평양": "PAC", "한화": "HAN",
    "해태": "HAE", "현대": "HYU", "kt": "KT",
}
ROLE = {"타자": "HITTER", "투수": "PITCHER"}

# 엑셀 열 이름 → data_pitch_type.pitch_code (엑셀 '서클 체인지업'은 띄어쓰기가 DB 와 다르다)
PITCH_CODE = {
    "포심": "FOUR_SEAM", "투심": "TWO_SEAM", "체인지업": "CHANGEUP",
    "서클 체인지업": "CIRCLE_CHANGEUP", "슬라이더": "SLIDER", "커브": "CURVE",
    "포크": "FORKBALL", "커터": "CUTTER", "싱커": "SINKER", "스플리터": "SPLITTER",
}

# data_player_card 와 다른 네임스페이스를 쓴다 — 카드 id 자체가 이미 이름 기반 v5 라
# 여기서 재활용하면 값이 같은 두 id 가 서로 다른 의미(카드 vs 구종행)로 충돌한다.
NAMESPACE_PITCH = uuid.uuid5(uuid.NAMESPACE_URL, "https://compyafun.com/data/player-card-pitch")

BATCH = 300  # INSERT 한 문장에 담을 행 수


def pitch_id(card_id: str, pitch_code: str) -> str:
    """card_id + pitch_code 기반 UUID v5. 재실행해도 같은 값이 나온다.

    legend_pitch 는 손으로 한 번 채운 자료라 임의 v4 로도 충분했지만, 이 스크립트는
    정본 엑셀이 갱신될 때마다 다시 돌 수 있어 카드 id 처럼 재현 가능한 v5 가 맞다.
    """
    return str(uuid.uuid5(NAMESPACE_PITCH, f"card_pitch:{card_id}:{pitch_code}"))


def team_code(raw: str) -> str:
    raw = raw.strip()
    return TEAM.get(raw, raw)


def load_card_ids() -> dict[tuple[str, str, str, str], str]:
    """(team_code, season_year, player_name, player_role) -> card_id"""
    keymap: dict[tuple[str, str, str, str], str] = {}
    with io.open(CARD_DUMP, encoding="utf-8") as f:
        for line in f:
            parts = line.rstrip("\n").split("\t")
            if len(parts) < 7:
                continue
            card_id, name, team, year, _pos, _sub, role = parts[:7]
            keymap[(team, year, name, role)] = card_id
    return keymap


def load_workbook():
    return openpyxl.load_workbook(XLSX, read_only=True, data_only=True)


def build_stat_rows(wb, keymap) -> tuple[list[tuple], list[str]]:
    """(카드id, stat1..5) 목록과, 못 지은 행 설명 목록."""
    rows: list[tuple] = []
    misses: list[str] = []
    sheets = (("일반_타자", "타자"), ("일반_투수", "투수"))
    for sheet_name, role_label in sheets:
        for r in wb[sheet_name].iter_rows(min_row=3, values_only=True):
            if not r or not r[0]:
                continue
            team = team_code(str(r[2]))
            year = str(int(float(r[3])))
            name = str(r[1]).strip()
            role = ROLE[role_label]
            key = (team, year, name, role)
            card_id = keymap.get(key)
            if card_id is None:
                misses.append(f"[{sheet_name}] {team} {year} {name} {role} — card_id 없음")
                continue
            stats = tuple(int(v) for v in r[8:13])
            rows.append((card_id, *stats))
    return rows, misses


def build_pitch_rows(wb, keymap) -> tuple[list[tuple], list[str]]:
    """(id, 카드id, pitch_code, grade) 목록과, 못 지은 행 설명 목록. '-' 는 만들지 않는다."""
    rows: list[tuple] = []
    misses: list[str] = []
    cols = list(PITCH_CODE.items())  # [(엑셀열이름, pitch_code), ...] — 열 순서대로
    for r in wb["투수_구종등급"].iter_rows(min_row=3, values_only=True):
        if not r or not r[0]:
            continue
        team = team_code(str(r[2]))
        year = str(int(float(r[3])))
        name = str(r[1]).strip()
        role = ROLE[str(r[6]).strip()]
        key = (team, year, name, role)
        card_id = keymap.get(key)
        if card_id is None:
            misses.append(f"[투수_구종등급] {team} {year} {name} {role} — card_id 없음")
            continue
        for i, (_col_name, pitch_code) in enumerate(cols):
            grade = r[8 + i]
            grade = str(grade).strip() if grade is not None else "-"
            if grade == "-":
                continue
            rows.append((pitch_id(card_id, pitch_code), card_id, pitch_code, grade))
    return rows, misses


def write_batched_insert(f, table: str, columns: str, rows: list[tuple], row_fmt) -> None:
    for i in range(0, len(rows), BATCH):
        chunk = rows[i:i + BATCH]
        f.write(f"INSERT INTO {table} ({columns})\nVALUES\n")
        f.write(",\n".join(row_fmt(r) for r in chunk))
        f.write(";\n\n")


def main() -> None:
    keymap = load_card_ids()
    wb = load_workbook()

    stat_rows, stat_misses = build_stat_rows(wb, keymap)
    pitch_rows, pitch_misses = build_pitch_rows(wb, keymap)

    misses = stat_misses + pitch_misses
    if misses:
        raise SystemExit(
            f"card_id 짝을 못 지은 행이 {len(misses)}건 있다 — 파일을 만들지 않는다.\n" +
            "\n".join(misses[:50]) +
            ("\n... (이하 생략)" if len(misses) > 50 else "")
        )

    hitter_count = sum(1 for _ in wb["일반_타자"].iter_rows(min_row=3, values_only=True) if _ and _[0])
    pitcher_count = sum(1 for _ in wb["일반_투수"].iter_rows(min_row=3, values_only=True) if _ and _[0])

    with io.open(OUT, "w", encoding="utf-8", newline="\n") as f:
        f.write(f"""-- =====================================================================
-- 노말 카드 스탯 + 구종등급 적재
--
-- 정본: {XLSX}
--   게임 화면 12,000회 캡처, 3회 교차검수. 일반_타자 {hitter_count}건 + 일반_투수
--   {pitcher_count}건 = 스탯 {len(stat_rows)}건. 에픽은 스탯이 비어 있어 범위 밖.
--
-- card_id 대응: 운영 DB 덤프(prod2.tsv)를 (구단, 연도, 이름, 역할)로 조회했다.
--   전 행이 짝을 지었다({len(stat_rows)}건) — 못 지은 행이 하나라도 있었으면 이
--   파일 자체가 생성되지 않는다(scripts/gen_card_stat_sql.py 의 검증 로직).
--
-- 구종등급: '-'(구종 없음)는 행을 만들지 않는다. 실측 결과 {len(pitch_rows)}건
--   (엑셀 투수_구종등급 시트 10개 구종 칸 중 '-' 아닌 칸 전체 — 이 스크립트가 직접 셌다).
--
-- 이 파일은 scripts/gen_card_stat_sql.py 가 만든다. 손으로 고치지 마라.
-- =====================================================================

SET NAMES utf8mb4;
USE compyafun;


-- 1) 노말 카드 태생 스탯 ({len(stat_rows)}건) -------------------------------------
""")
        write_batched_insert(
            f, "data_player_card_stat", "card_id, stat1, stat2, stat3, stat4, stat5",
            stat_rows,
            lambda r: f"('{r[0]}', {r[1]}, {r[2]}, {r[3]}, {r[4]}, {r[5]})",
        )

        f.write(f"-- 2) 노말 카드(투수) 보유 구종 ({len(pitch_rows)}건) -----------------------------\n")
        write_batched_insert(
            f, "data_player_card_pitch", "id, card_id, pitch_code, pitch_grade",
            pitch_rows,
            lambda r: f"('{r[0]}', '{r[1]}', '{r[2]}', '{r[3]}')",
        )

        f.write(f"""
-- 3) 적재 후 검증 --------------------------------------------------------

--    스탯 행 수 — {len(stat_rows)}건이어야 한다 (타자 {hitter_count} + 투수 {pitcher_count})
SELECT COUNT(*) AS 스탯_행수 FROM data_player_card_stat;

--    구종 행 수 — {len(pitch_rows)}건이어야 한다 ('-' 제외 실측치)
SELECT COUNT(*) AS 구종_행수 FROM data_player_card_pitch;

--    타자 카드에 구종 행이 붙어 있으면 안 된다 (0건이어야 정상)
SELECT COUNT(*) AS 타자인데_구종있음
FROM data_player_card_pitch p
JOIN data_player_card c ON c.id = p.card_id
WHERE c.player_role = 'HITTER';

--    카드와 짝이 안 맞는 스탯 행 (0건이어야 정상 — FK 로도 막히지만 이중 확인)
SELECT COUNT(*) AS 스탯_고아행
FROM data_player_card_stat s
LEFT JOIN data_player_card c ON c.id = s.card_id
WHERE c.id IS NULL;

--    스탯이 없는 노말 카드 (0건이어야 정상 — {len(stat_rows)}건 전량 적재 확인)
SELECT COUNT(*) AS 스탯_누락_카드
FROM data_player_card c
LEFT JOIN data_player_card_stat s ON s.card_id = c.id
WHERE c.card_type = 'NORMAL' AND s.card_id IS NULL;
""")

    print(f"{OUT} — stat {len(stat_rows)}건, pitch {len(pitch_rows)}건")


if __name__ == "__main__":
    main()
